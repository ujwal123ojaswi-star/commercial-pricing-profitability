import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
rng=np.random.default_rng(2026)

# ---------- synthetic master data ----------
fam={"FCC":(3000,4500,.62,.70),"Hydroprocessing":(8000,15000,.55,.65),"Specialty":(20000,40000,.45,.58)}
prods=[]
for f,(lo,hi,cl,ch) in fam.items():
    for i in range(4):
        lp=round(rng.uniform(lo,hi),-1); prods.append(dict(Product=f"{f[:3].upper()}-{i+1:02d}",Family=f,List25=lp,Cost25=round(lp*rng.uniform(cl,ch),-1)))
P=pd.DataFrame(prods)
segs=["Refining","Petrochemical","Specialty Chemicals"]
cust=[]
for i in range(40):
    tier=rng.choice(["Strategic","Core","Transactional"],p=[.2,.4,.4])
    d={"Strategic":(.10,.18),"Core":(.05,.11),"Transactional":(.02,.07)}[tier]
    cust.append(dict(Customer=f"C{i+1:03d}",Segment=segs[i%3] if i%5 else "Refining",Tier=tier,
      Disc25=round(rng.uniform(*d),3),Rebate=round(rng.choice([0,.01,.02,.03,.05],p=[.3,.25,.2,.15,.1]),3),
      Freight=round(rng.uniform(80,420),-1),Terms=int(rng.choice([30,60,90],p=[.4,.4,.2])),
      CTS=round(rng.uniform(*{"Strategic":(700000,1400000),"Core":(300000,700000),"Transactional":(150000,420000)}[tier]),-3),
      DiscDrift=round(max(0,rng.normal(.015,.02)) if rng.random()<.45 else 0,3),
      VolChg=round(rng.normal(.0,.12),3)))
C=pd.DataFrame(cust)
rows=[]
for _,c in C.iterrows():
    fams=list(fam) if c.Segment!="Refining" else ["FCC","Hydroprocessing","Hydroprocessing","Specialty"]
    pick=P[P.Family.isin(set(fams))].sample(int(rng.integers(2,5)),random_state=int(rng.integers(1e6)))
    base=rng.uniform(3,24)*{"Strategic":3,"Core":1.5,"Transactional":.7}[c.Tier]
    for _,p in pick.iterrows():
        t=base/(p.List25/4000)**.6
        shift=1.12 if p.Family=="Specialty" else .97
        for y in (2025,2026):
            for m in range(1,13):
                tons=t*rng.uniform(.7,1.3)*((1+c.VolChg)*shift if y==2026 else 1)
                rows.append(dict(Year=str(y),Month=m,Customer=c.Customer,Product=p.Product,Tons=round(tons,1)))
T=pd.DataFrame(rows)

# ---------- workbook ----------
wb=Workbook(); F="Arial"
H=Font(name=F,bold=True,color="FFFFFF"); HF=PatternFill("solid",fgColor="1F3864")
BL=Font(name=F,color="0000FF"); BK=Font(name=F); GR=Font(name=F,color="008000"); B=Font(name=F,bold=True)
Y=PatternFill("solid",fgColor="FFFF00")
USD='$#,##0;($#,##0);-'; USD2='$#,##0.00;($#,##0.00);-'; PCT='0.0%;(0.0%);-'; NUM='#,##0.0;(#,##0.0);-'
def hdr(ws,r,cols):
    for j,v in enumerate(cols,1):
        c=ws.cell(r,j,v); c.font=H; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical="center")
def widths(ws,w):
    for j,x in enumerate(w,1): ws.column_dimensions[L(j)].width=x

# Assumptions
A=wb.active; A.title="Assumptions"
A["A1"]="Commercial Pricing & Profitability Analytics: Assumptions"; A["A1"].font=Font(name=F,bold=True,size=13)
A["A2"]="Synthetic B2B catalyst dataset (illustrative only). Blue = input, black = formula, green = cross-sheet link. Yellow = key levers."; A["A2"].font=Font(name=F,italic=True)
inp=[("FY2026 announced list price increase",0.05,PCT),("FY2026 standard cost inflation",0.03,PCT),
     ("Cost of capital (payment-terms cost)",0.08,PCT),("Target customer net margin (renewal scenario)",0.25,PCT)]
for i,(k,v,f) in enumerate(inp,4):
    A.cell(i,1,k).font=BK; c=A.cell(i,2,v); c.font=BL; c.number_format=f; c.fill=Y
A["C4"]="Assumption: uniform list increase across all grades"; A["C7"]="Floor for contract renewal pricing"
for r in (4,7): A.cell(r,3).font=Font(name=F,italic=True)
widths(A,[48,12,50])

# Products
Pw=wb.create_sheet("Products"); hdr(Pw,1,["Product","Family","List Price FY25 ($/t)","Std Cost FY25 ($/t)","List Price FY26 ($/t)","Std Cost FY26 ($/t)"])
for i,p in enumerate(prods,2):
    Pw.cell(i,1,p["Product"]).font=BK; Pw.cell(i,2,p["Family"]).font=BK
    for j,k in ((3,"List25"),(4,"Cost25")): c=Pw.cell(i,j,p[k]); c.font=BL; c.number_format=USD
    c=Pw.cell(i,5,f"=C{i}*(1+Assumptions!$B$4)"); c.number_format=USD; c.font=GR
    c=Pw.cell(i,6,f"=D{i}*(1+Assumptions!$B$5)"); c.number_format=USD; c.font=GR
NP=len(prods)+1; widths(Pw,[11,17,14,14,14,14])

# Customers
Cw=wb.create_sheet("Customers")
hdr(Cw,1,["Customer","Segment","Tier","On-Invoice Discount FY25","Discount Drift FY26","On-Invoice Discount FY26","Volume Rebate","Freight ($/t)","Payment Terms (days)","Annual Cost-to-Serve ($)"])
for i,c in enumerate(cust,2):
    for j,k in ((1,"Customer"),(2,"Segment"),(3,"Tier")): Cw.cell(i,j,c[k]).font=BK
    for j,k,f in ((4,"Disc25",PCT),(5,"DiscDrift",PCT),(7,"Rebate",PCT),(8,"Freight",USD),(9,"Terms","0"),(10,"CTS",USD)):
        x=Cw.cell(i,j,c[k]); x.font=BL; x.number_format=f
    x=Cw.cell(i,6,f"=D{i}+E{i}"); x.number_format=PCT
NC=len(cust)+1; widths(Cw,[10,19,13,12,11,12,10,10,10,14])

# Transactions with row formulas
D=wb.create_sheet("Transactions")
cols=["Year","Month","Customer","Product","Family","Tons","List Price ($/t)","On-Inv Disc","Gross Revenue ($)","Invoice Revenue ($)","Rebates ($)","Freight ($)","Terms Cost ($)","Pocket Revenue ($)","Std COGS ($)","Pocket Margin ($)"]
hdr(D,1,cols)
for i,t in enumerate(T.itertuples(index=False),2):
    D.cell(i,1,t.Year); D.cell(i,2,t.Month); D.cell(i,3,t.Customer); D.cell(i,4,t.Product)
    x=D.cell(i,6,t.Tons); x.font=BL; x.number_format=NUM
    pm=f"Products!$A$2:$A${NP}"; cm=f"Customers!$A$2:$A${NC}"
    D.cell(i,5,f"=INDEX(Products!$B$2:$B${NP},MATCH(D{i},{pm},0))")
    D.cell(i,7,f'=IF(A{i}="2025",INDEX(Products!$C$2:$C${NP},MATCH(D{i},{pm},0)),INDEX(Products!$E$2:$E${NP},MATCH(D{i},{pm},0)))').number_format=USD
    D.cell(i,8,f'=IF(A{i}="2025",INDEX(Customers!$D$2:$D${NC},MATCH(C{i},{cm},0)),INDEX(Customers!$F$2:$F${NC},MATCH(C{i},{cm},0)))').number_format=PCT
    D.cell(i,9,f"=F{i}*G{i}").number_format=USD
    D.cell(i,10,f"=I{i}*(1-H{i})").number_format=USD
    D.cell(i,11,f"=J{i}*INDEX(Customers!$G$2:$G${NC},MATCH(C{i},{cm},0))").number_format=USD
    D.cell(i,12,f"=F{i}*INDEX(Customers!$H$2:$H${NC},MATCH(C{i},{cm},0))").number_format=USD
    D.cell(i,13,f"=J{i}*INDEX(Customers!$I$2:$I${NC},MATCH(C{i},{cm},0))/365*Assumptions!$B$6").number_format=USD
    D.cell(i,14,f"=J{i}-K{i}-L{i}-M{i}").number_format=USD
    D.cell(i,15,f'=F{i}*IF(A{i}="2025",INDEX(Products!$D$2:$D${NP},MATCH(D{i},{pm},0)),INDEX(Products!$F$2:$F${NP},MATCH(D{i},{pm},0)))').number_format=USD
    D.cell(i,16,f"=N{i}-O{i}").number_format=USD
NT=len(T)+1; widths(D,[7,7,10,10,16,9]+[14]*10); D.freeze_panes="A2"
rng_=lambda col:f"Transactions!${col}$2:${col}${NT}"

# Price Waterfall
W=wb.create_sheet("Price Waterfall")
W["A1"]="Pocket Price Waterfall ($/ton, all customers)"; W["A1"].font=Font(name=F,bold=True,size=13)
hdr(W,3,["Step","FY2025 ($/t)","FY2026 ($/t)","FY2026 % of List"])
steps=[("List price","I"),("On-invoice discount","I-J"),("Invoice price","J"),("Volume rebates","K"),("Freight","L"),("Payment-terms cost","M"),("Pocket price","N")]
for r,(name,col) in enumerate(steps,4):
    W.cell(r,1,name).font=B if name in("List price","Invoice price","Pocket price") else BK
    for j,y in ((2,"2025"),(3,"2026")):
        tons=f'SUMIFS({rng_("F")},{rng_("A")},"{y}")'
        if "-" in col:
            a,b=col.split("-"); f=f'=-(SUMIFS({rng_(a)},{rng_("A")},"{y}")-SUMIFS({rng_(b)},{rng_("A")},"{y}"))/{tons}'
        elif name in("Volume rebates","Freight","Payment-terms cost"):
            f=f'=-SUMIFS({rng_(col)},{rng_("A")},"{y}")/{tons}'
        else: f=f'=SUMIFS({rng_(col)},{rng_("A")},"{y}")/{tons}'
        W.cell(r,j,f).number_format=USD2
    W.cell(r,4,f"=C{r}/$C$4").number_format=PCT
W["A12"]="Price leakage, list to pocket"; W["A12"].font=B
W["B12"]="=1-B10/B4"; W["C12"]="=1-C10/C4"
for c in("B12","C12"): W[c].number_format=PCT
W["A14"]="Price Realization (FY2026 vs FY2025)"; W["A14"].font=Font(name=F,bold=True,size=12)
hdr(W,15,["Measure","Change"])
W["A16"]="Announced list increase"; W["B16"]="=Assumptions!B4"; W["B16"].font=GR
W["A17"]="Realized invoice price change (same-mix basis)"
W["A18"]="Realized pocket price change (same-mix basis)"
W["A19"]="Realization rate (pocket change / announced)"; W["B19"]="=B18/B16"
for c in("B16","B17","B18","B19"): W[c].number_format=PCT
widths(W,[46,16,16,16])

# Product PVM
V=wb.create_sheet("PVM Bridge")
V["A1"]="Price-Volume-Mix Bridge: Pocket Revenue FY2025 to FY2026, by product"; V["A1"].font=Font(name=F,bold=True,size=13)
hdr(V,3,["Product","Family","Tons FY25","Tons FY26","Pocket $/t FY25","Pocket $/t FY26","Pocket Rev FY25","Pocket Rev FY26","FY26 Tons at FY25 Price","Price Effect","Pocket Margin FY25","Pocket Margin FY26"])
r0=4
for k,p in enumerate(prods):
    r=r0+k; V.cell(r,1,p["Product"]); V.cell(r,2,p["Family"])
    for j,y in ((3,"2025"),(4,"2026")):
        V.cell(r,j,f'=SUMIFS({rng_("F")},{rng_("D")},$A{r},{rng_("A")},"{y}")').number_format=NUM
    for j,y in ((7,"2025"),(8,"2026")):
        V.cell(r,j,f'=SUMIFS({rng_("N")},{rng_("D")},$A{r},{rng_("A")},"{y}")').number_format=USD
    for j,y in ((11,"2025"),(12,"2026")):
        V.cell(r,j,f'=SUMIFS({rng_("P")},{rng_("D")},$A{r},{rng_("A")},"{y}")').number_format=USD
    V.cell(r,5,f"=IFERROR(G{r}/C{r},0)").number_format=USD2
    V.cell(r,6,f"=IFERROR(H{r}/D{r},0)").number_format=USD2
    V.cell(r,9,f"=D{r}*E{r}").number_format=USD
    V.cell(r,10,f"=D{r}*(F{r}-E{r})").number_format=USD
rT=r0+len(prods)
V.cell(rT,1,"Total").font=B
for j in (3,4,7,8,9,10,11,12):
    c=V.cell(rT,j,f"=SUM({L(j)}{r0}:{L(j)}{rT-1})"); c.font=B; c.number_format=NUM if j in(3,4) else USD
b=rT+2
V.cell(b,1,"Bridge").font=Font(name=F,bold=True,size=12); hdr(V,b+1,["Component","Pocket Revenue ($)","% of FY25"])
items=[("FY2025 pocket revenue",f"=G{rT}"),("Volume effect",f"=(D{rT}-C{rT})*(G{rT}/C{rT})"),
       ("Mix effect",f"=I{rT}-D{rT}*(G{rT}/C{rT})"),("Price effect",f"=J{rT}"),("FY2026 pocket revenue",f"=SUM(B{b+2}:B{b+5})"),
       ("Check vs actual FY2026 (should be 0)",f"=B{b+6}-H{rT}")]
for k,(n,f) in enumerate(items):
    r=b+2+k; V.cell(r,1,n).font=B if k in(0,4) else BK
    V.cell(r,2,f).number_format=USD; V.cell(r,3,f"=B{r}/$B${b+2}").number_format=PCT
mrow=b+9
V.cell(mrow,1,"Pocket margin % FY25 / FY26").font=B
V.cell(mrow,2,f"=K{rT}/G{rT}").number_format=PCT; V.cell(mrow,3,f"=L{rT}/H{rT}").number_format=PCT
widths(V,[34,16,11,11,13,13,15,15,16,14,15,15])
# realization (same-mix): FY26 tons at FY25 price vs FY26 actual
W["B17"]=(f'=SUMIFS({rng_("J")},{rng_("A")},"2026")/SUMPRODUCT(\'PVM Bridge\'!D{r0}:D{rT-1},'
          f'\'PVM Bridge\'!E{r0}:E{rT-1}+0)*0+0')
# simpler & exact: invoice same-mix needs FY25 invoice $/t per product; compute helper cols
V.cell(3,13,"Invoice $/t FY25").font=H; V.cell(3,13).fill=HF; V.cell(3,14,"Invoice $/t FY26").font=H; V.cell(3,14).fill=HF
for k in range(len(prods)):
    r=r0+k
    for j,y,tc in ((13,"2025","C"),(14,"2026","D")):
        V.cell(r,j,f'=IFERROR(SUMIFS({rng_("J")},{rng_("D")},$A{r},{rng_("A")},"{y}")/{tc}{r},0)').number_format=USD2
V.column_dimensions["M"].width=14; V.column_dimensions["N"].width=14
W["B17"]=f"=SUMPRODUCT('PVM Bridge'!D{r0}:D{rT-1},'PVM Bridge'!N{r0}:N{rT-1})/SUMPRODUCT('PVM Bridge'!D{r0}:D{rT-1},'PVM Bridge'!M{r0}:M{rT-1})-1"
W["B18"]=f"='PVM Bridge'!H{rT}/'PVM Bridge'!I{rT}-1"
for c in ("B17","B18"): W[c].font=GR

# Customer profitability (rows ordered by python estimate later -> recompute order after recalc)
U=wb.create_sheet("Customer Profitability")
U["A1"]="Customer Profitability FY2026 (pocket margin less cost-to-serve)"; U["A1"].font=Font(name=F,bold=True,size=13)
hdr(U,3,["Customer","Segment","Tier","Pocket Revenue ($)","Pocket Margin ($)","Cost-to-Serve ($)","Customer Net Profit ($)","Net Margin %","Rank","Cumulative Profit % (whale curve)","Price Uplift to Target Margin"])
wb.save("tmp.xlsx")
pd.to_pickle((C,P,T,prods,cust,NC,NT,r0,rT,b),"state.pkl")
