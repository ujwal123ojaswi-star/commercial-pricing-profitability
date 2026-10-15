import pandas as pd, numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font
C,P,T,prods,cust,NC,NT,r0,rT,b=pd.read_pickle("state.pkl")
# python replica to determine sort order
df=T.merge(P,on="Product").merge(C,on="Customer")
y26=df.Year=="2026"
df["List"]=np.where(y26,df.List25*1.05,df.List25); df["Cost"]=np.where(y26,df.Cost25*1.03,df.Cost25)
df["Disc"]=np.where(y26,df.Disc25+df.DiscDrift,df.Disc25)
df["Inv"]=df.Tons*df.List*(1-df.Disc)
df["Pocket"]=df.Inv*(1-df.Rebate)-df.Tons*df.Freight-df.Inv*df.Terms/365*.08
df["PM"]=df.Pocket-df.Tons*df.Cost
g=df[y26].groupby("Customer")[["Pocket","PM"]].sum().join(C.set_index("Customer"))
g["Net"]=g.PM-g.CTS; order=g.sort_values("Net",ascending=False).index.tolist()
wb=load_workbook("tmp.xlsx"); U=wb["Customer Profitability"]
F="Arial"; USD='$#,##0;($#,##0);-'; PCT='0.0%;(0.0%);-'
R=lambda c:f"Transactions!${c}$2:${c}${NT}"; cm=f"Customers!$A$2:$A${NC}"
s=4; e=s+len(order)-1
for k,cu in enumerate(order):
    r=s+k; U.cell(r,1,cu)
    U.cell(r,2,f"=INDEX(Customers!$B$2:$B${NC},MATCH(A{r},{cm},0))")
    U.cell(r,3,f"=INDEX(Customers!$C$2:$C${NC},MATCH(A{r},{cm},0))")
    U.cell(r,4,f'=SUMIFS({R("N")},{R("C")},$A{r},{R("A")},"2026")').number_format=USD
    U.cell(r,5,f'=SUMIFS({R("P")},{R("C")},$A{r},{R("A")},"2026")').number_format=USD
    U.cell(r,6,f"=INDEX(Customers!$J$2:$J${NC},MATCH(A{r},{cm},0))").number_format=USD
    U.cell(r,7,f"=E{r}-F{r}").number_format=USD
    U.cell(r,8,f"=IFERROR(G{r}/D{r},0)").number_format=PCT
    U.cell(r,9,f"=RANK(G{r},$G${s}:$G${e})")
    U.cell(r,10,f"=SUM($G${s}:G{r})/SUM($G${s}:$G${e})").number_format=PCT
    U.cell(r,11,f"=MAX(0,(Assumptions!$B$7*D{r}-G{r})/(1-Assumptions!$B$7))/D{r}").number_format=PCT
t=e+2
lab=[("Total customers",f"=COUNTA(A{s}:A{e})","0"),
     ("Customers with negative net profit",f'=COUNTIF(G{s}:G{e},"<0")',"0"),
     ("Customers below target net margin",f'=COUNTIF(H{s}:H{e},"<"&Assumptions!B7)',"0"),
     ("Peak cumulative profit (whale curve)",f"=MAX(J{s}:J{e})",PCT),
     ("Customers needed to reach peak profit",f"=MATCH(MAX(J{s}:J{e}),J{s}:J{e},0)","0"),
     ("Profit erased by loss-making customers",f'=-SUMIF(G{s}:G{e},"<0")',USD),
     ("Total customer net profit",f"=SUM(G{s}:G{e})",USD),
     ("Pocket revenue at risk (below target)",f'=SUMIF(H{s}:H{e},"<"&Assumptions!B7,D{s}:D{e})',USD),
     ("Avg price uplift needed (below-target customers)",f'=SUMPRODUCT((H{s}:H{e}<Assumptions!B7)*K{s}:K{e}*D{s}:D{e})/H{t+7}',PCT)]
for k,(n,f,nf) in enumerate(lab):
    r=t+k; U.cell(r,1,n).font=Font(name=F,bold=True); c=U.cell(r,8,f); c.number_format=nf
# fix last formula ref to revenue-at-risk cell
U.cell(t+8,8).value=f'=IFERROR(SUMPRODUCT((H{s}:H{e}<Assumptions!B7)*K{s}:K{e}*D{s}:D{e})/H{t+7},0)'
for j,w in enumerate([10,19,13,15,15,14,16,11,7,15,14],1): U.column_dimensions[chr(64+j)].width=w
U.freeze_panes="A4"
# Summary dashboard sheet first
S=wb.create_sheet("Summary",0)
S["A1"]="Commercial Pricing & Profitability Analytics: FY2026 Summary"; S["A1"].font=Font(name=F,bold=True,size=14)
S["A2"]="Synthetic B2B specialty-catalyst dataset: 40 customers, 12 products, 3 families, 24 months. Illustrative only."; S["A2"].font=Font(name=F,italic=True)
G=Font(name=F,color="008000")
k=[("Announced list increase","='Price Waterfall'!B16",PCT),("Realized invoice price change","='Price Waterfall'!B17",PCT),
   ("Realized pocket price change","='Price Waterfall'!B18",PCT),("Realization rate","='Price Waterfall'!B19",PCT),
   ("Price leakage list to pocket FY26","='Price Waterfall'!C12",PCT),
   ("Pocket revenue FY25",f"='PVM Bridge'!B{b+2}",USD),("Volume effect",f"='PVM Bridge'!B{b+3}",USD),
   ("Mix effect",f"='PVM Bridge'!B{b+4}",USD),("Price effect",f"='PVM Bridge'!B{b+5}",USD),("Pocket revenue FY26",f"='PVM Bridge'!B{b+6}",USD),
   ("Pocket margin % FY25",f"='PVM Bridge'!B{b+9}",PCT),("Pocket margin % FY26",f"='PVM Bridge'!C{b+9}",PCT),
   ("Loss-making customers",f"='Customer Profitability'!H{t+1}","0"),("Customers below target margin",f"='Customer Profitability'!H{t+2}","0"),
   ("Peak cumulative profit (whale curve)",f"='Customer Profitability'!H{t+3}",PCT),("Customers to reach peak",f"='Customer Profitability'!H{t+4}","0"),
   ("Profit erased by loss-makers",f"='Customer Profitability'!H{t+5}",USD),("Avg uplift needed, below-target customers",f"='Customer Profitability'!H{t+8}",PCT)]
for i,(n,f,nf) in enumerate(k,4):
    S.cell(i,1,n).font=Font(name=F); c=S.cell(i,2,f); c.font=G; c.number_format=nf
S.column_dimensions["A"].width=42; S.column_dimensions["B"].width=18
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font and c.font.name!=F: c.font=Font(name=F,bold=c.font.bold,italic=c.font.italic,color=c.font.color,size=c.font.size)
wb.save("Commercial_Pricing_Profitability.xlsx")
